# 调试导入问题 - 模拟实际使用场景
import os
import sys
from openpyxl import Workbook

# 添加当前目录到路径
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from config import OUTPUT_DIR
from excel_importer import ExcelImporter
from main import ImportMode, PartyFeeSystem


def create_test_file_with_formulas():
    """创建一个包含公式的测试文件（模拟用户从系统导出的文件）"""
    print("\n" + "=" * 60)
    print("创建测试文件（包含公式）")
    print("=" * 60)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "明细表"
    
    # 第1行：标题
    ws['A1'] = "党费收缴明细表"
    ws.merge_cells('A1:R1')
    
    # 第2行：分组表头
    ws['E2'] = "工资项目"
    ws.merge_cells('E2:H2')
    ws['I2'] = "扣款项目"
    ws.merge_cells('I2:P2')
    
    # 第3行：列名
    headers = [
        '支部序号', '所属党支部', '序号', '姓 名',
        '岗位工资', '薪级工资', '高定工资', '基础性绩效',
        '住房公积金', '医疗保险', '养老保险', '职业年金',
        '大额医疗', '失业保险', '个人所得税',
        '缴费基数', '每月应缴党费'
    ]
    
    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=3, column=col_idx, value=header)
    
    # 第4行开始：数据（使用公式）
    test_data = [
        (1, '教工第一党支部', 1, '张三', 5000, 3000, 0, 2000, 800, 200, 400, 100, 50, 50, 0),
        (1, '教工第一党支部', 2, '李四', 6000, 4000, 0, 2500, 900, 250, 500, 120, 60, 60, 0),
    ]
    
    for row_idx, row_data in enumerate(test_data, 4):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
        
        # 缴费基数公式：工资项目总和 - 扣款项目总和
        # E列到H列是工资，I列到O列是扣款
        ws.cell(row=row_idx, column=16, value=f'=SUM(E{row_idx}:H{row_idx})-SUM(I{row_idx}:O{row_idx})')
        # 党费公式（简化）
        ws.cell(row=row_idx, column=17, value=f'=P{row_idx}*0.015')
    
    # 保存文件
    test_file_path = os.path.join(OUTPUT_DIR, '测试_包含公式.xlsx')
    wb.save(test_file_path)
    
    print(f"[OK] 已创建测试文件: {test_file_path}")
    print(f"     注意: 这个文件包含公式，如果从未用Excel打开过，data_only=True会读取到None")
    
    return test_file_path


def test_import_with_data_only():
    """测试使用 data_only=True 导入文件"""
    print("\n" + "=" * 60)
    print("测试1：导入包含公式的文件（data_only=True的影响）")
    print("=" * 60)
    
    file_path = create_test_file_with_formulas()
    
    importer = ExcelImporter()
    
    # 测试1：使用默认的 data_only=True
    print("\n--- 使用 data_only=True 加载 ---")
    from openpyxl import load_workbook
    
    wb1 = load_workbook(file_path, data_only=True)
    ws1 = wb1.active
    
    print("\n读取数据行（data_only=True）:")
    for row in range(4, 6):
        name = ws1.cell(row=row, column=4).value
        payment_base = ws1.cell(row=row, column=16).value
        monthly_fee = ws1.cell(row=row, column=17).value
        print(f"  行{row}: 姓名={name}, 缴费基数={payment_base}, 党费={monthly_fee}")
    
    wb1.close()
    
    # 测试2：不使用 data_only
    print("\n--- 不使用 data_only 加载 ---")
    wb2 = load_workbook(file_path, data_only=False)
    ws2 = wb2.active
    
    print("\n读取数据行（data_only=False）:")
    for row in range(4, 6):
        name = ws2.cell(row=row, column=4).value
        payment_base = ws2.cell(row=row, column=16).value
        monthly_fee = ws2.cell(row=row, column=17).value
        print(f"  行{row}: 姓名={name}, 缴费基数={payment_base}, 党费={monthly_fee}")
    
    wb2.close()
    
    # 测试3：实际导入
    print("\n--- 实际导入测试 ---")
    result = importer.import_from_file(file_path, sheet_type='detail')
    
    print(f"\n导入结果:")
    print(f"  成功: {result.success}")
    print(f"  支部数量: {result.total_branches}")
    print(f"  党员数量: {result.total_members}")
    
    if result.errors:
        print(f"\n错误:")
        for e in result.errors:
            print(f"  - {e}")
    
    if result.warnings:
        print(f"\n警告:")
        for w in result.warnings:
            print(f"  - {w}")
    
    return result


def test_load_data_exit_mechanism():
    """测试加载数据的退出机制"""
    print("\n" + "=" * 60)
    print("测试2：分析加载数据的退出机制")
    print("=" * 60)
    
    # 查看 _auto_load_data 方法
    print("\n分析 _auto_load_data 方法:")
    print("  提供的选项:")
    print("    1,2,...N: 加载对应月份的数据")
    print("    0: 开始新的月份（不加载历史数据）")
    print("  问题: 没有提供直接退出系统的选项！")
    print("  如果用户想退出，只能先选择0或其他选项，然后在主菜单选择0退出")
    
    # 查看 load_data 方法
    print("\n分析 load_data 方法:")
    print("  该方法会询问年份和月份")
    print("  问题: 没有提供退出机制！")
    print("  用户输入 'q' 或 'quit' 不会被识别，会触发 ValueError")
    
    # 查看主菜单的加载数据选项
    print("\n主菜单中的'加载数据'选项:")
    print("  选择10后，会调用 load_data() 方法")
    print("  load_data() 方法中:")
    print("    year = self.input_int('请输入年份', self.current_year)")
    print("    month = self.input_int('请输入月份', self.current_month)")
    print("  input_int 方法不支持退出！")


def analyze_input_functions():
    """分析输入函数"""
    print("\n" + "=" * 60)
    print("测试3：分析输入函数")
    print("=" * 60)
    
    # 查看 input_int 和 input_float 方法
    print("\ninput_int 方法:")
    print("  def input_int(self, prompt: str, default: int = 0) -> int:")
    print("      while True:")
    print("          try:")
    print("              value = input(f\"{prompt} (默认: {default}): \")")
    print("              if value.strip() == \"\":")
    print("                  return default")
    print("              return int(value)")
    print("          except ValueError:")
    print("              print(\"请输入有效的整数！\")")
    print("  问题: 不支持 'q' 或 'quit' 退出！")
    
    print("\ninput_float 方法同理，也不支持退出。")


def test_import_integration():
    """测试导入的完整流程（模拟用户实际操作）"""
    print("\n" + "=" * 60)
    print("测试4：模拟完整的导入流程")
    print("=" * 60)
    
    # 创建测试文件
    from test_multi_header_import import create_test_multi_header_file
    file_path = create_test_multi_header_file()
    
    print("\n--- 步骤1：导入数据 ---")
    importer = ExcelImporter()
    result = importer.import_from_file(file_path, sheet_type='detail')
    
    print(f"\n导入结果:")
    print(f"  成功: {result.success}")
    print(f"  支部数量: {result.total_branches}")
    print(f"  党员数量: {result.total_members}")
    
    if result.success:
        print("\n导入的数据详情:")
        for branch in result.branches:
            print(f"\n  支部: {branch.name}")
            for member in branch.members:
                print(f"    - {member.name}:")
                print(f"      岗位工资: {member.salary_info.position_salary}")
                print(f"      薪级工资: {member.salary_info.rank_salary}")
                print(f"      缴费基数: {member.payment_base}")
                print(f"      党费: {member.monthly_fee}")
    
    # 模拟导入后的数据处理
    print("\n--- 步骤2：模拟导入后的数据处理 ---")
    print("\n在 main.py 中，导入后会调用:")
    print("  - _import_append() : 追加模式")
    print("  - _import_replace() : 替换模式")
    print("  - _import_merge() : 合并模式")
    
    print("\n检查 _import_replace 方法:")
    print("  def _import_replace(self, new_branches: List[PartyBranch]):")
    print("      # 清除现有数据")
    print("      self.member_manager = MemberManager()")
    print("      # 添加导入的数据")
    print("      for branch in new_branches:")
    print("          self.member_manager.add_branch(branch)")
    print("  这个逻辑看起来是正确的...")
    
    print("\n检查 add_branch 方法:")
    print("  def add_branch(self, branch: PartyBranch) -> None:")
    print("      # 检查是否已存在同名支部")
    print("      for existing_branch in self.branches:")
    print("          if existing_branch.name == branch.name:")
    print("              raise ValueError(f\"党支部 '{branch.name}' 已存在\")")
    print("          if existing_branch.sequence == branch.sequence:")
    print("              raise ValueError(f\"支部序号 '{branch.sequence}' 已存在\")")
    print("      self.branches.append(branch)")
    print("      # 按支部序号排序")
    print("      self.branches.sort(key=lambda b: b.sequence)")
    print("  这个逻辑也是正确的...")
    
    return result


if __name__ == "__main__":
    print("\n" + "#" * 60)
    print("# 调试导入问题和退出机制问题")
    print("#" * 60)
    
    # 测试1：导入包含公式的文件
    test_import_with_data_only()
    
    # 测试2：分析退出机制
    test_load_data_exit_mechanism()
    
    # 测试3：分析输入函数
    analyze_input_functions()
    
    # 测试4：完整导入流程
    test_import_integration()
    
    print("\n" + "#" * 60)
    print("# 分析完成")
    print("#" * 60)
    
    print("\n总结发现的问题:")
    print("\n问题1：导入数据为空")
    print("  可能的原因:")
    print("  1. 使用 data_only=True 加载包含公式的文件，如果文件未被Excel打开过，公式结果为None")
    print("  2. 但测试显示姓名等纯数据应该能正常读取...")
    print("  3. 可能需要进一步检查实际用户的Excel文件结构")
    
    print("\n问题2：加载数据没有退出机制")
    print("  具体问题:")
    print("  1. _auto_load_data 方法没有提供退出系统的选项")
    print("  2. load_data 方法没有提供退出选项")
    print("  3. input_int 和 input_float 方法不支持 'q' 或 'quit' 退出")
