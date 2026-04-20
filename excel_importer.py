# Excel文件导入模块
import os
import re
from typing import List, Dict, Optional, Tuple, Any
from dataclasses import dataclass

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter

from data_models import PartyMember, PartyBranch, SalaryInfo, DeductionInfo
from fee_calculator import FeeCalculator


@dataclass
class ImportResult:
    """导入结果"""
    success: bool = False
    branches: List[PartyBranch] = None
    errors: List[str] = None
    warnings: List[str] = None
    total_members: int = 0
    total_branches: int = 0
    
    def __post_init__(self):
        if self.branches is None:
            self.branches = []
        if self.errors is None:
            self.errors = []
        if self.warnings is None:
            self.warnings = []


class ExcelImporter:
    """Excel文件导入器"""
    
    def __init__(self):
        """初始化导入器"""
        # 定义列名映射（支持多种可能的列名）
        self.column_mappings = {
            # 基本信息
            'branch_sequence': ['支部序号', '支部编号', '支部号'],
            'branch_name': ['所属党支部', '支部名称', '党支部', '所属支部'],
            'sequence': ['序号', '编号'],
            'name': ['姓名', '党员姓名'],
            
            # 工资信息
            'position_salary': ['岗位工资'],
            'rank_salary': ['薪级工资'],
            'fixed_salary': ['高定工资'],
            'basic_performance': ['基础性绩效', '绩效工资', '基础绩效'],
            
            # 扣款信息
            'housing_fund': ['住房公积金', '公积金', '住房公积'],
            'medical_insurance': ['医疗保险', '医保'],
            'pension_insurance': ['养老保险', '养老'],
            'occupational_annuity': ['职业年金', '年金'],
            'large_medical': ['大额医疗', '大额医保'],
            'unemployment_insurance': ['失业保险', '失业'],
            'personal_income_tax': ['个人所得税', '个税', '所得税'],
            
            # 计算结果
            'payment_base': ['缴费基数', '基数'],
            'monthly_fee': ['每月应缴党费', '月党费', '党费', '应缴党费']
        }
        
        # 反转映射，用于快速查找
        self.reverse_mapping = {}
        for key, values in self.column_mappings.items():
            for value in values:
                self.reverse_mapping[value] = key
    
    def _normalize_column_name(self, name: str) -> str:
        """
        规范化列名
        
        去除空格、特殊字符、中间空格等
        """
        if not name:
            return ""
        
        # 转换为字符串
        name = str(name)
        
        # 去除换行符和回车符
        name = name.replace('\n', '').replace('\r', '')
        
        # 去除括号中的内容，如"缴费基数（元）" -> "缴费基数"
        name = re.sub(r'[（(].*?[)）]', '', name)
        
        # 去除所有空格（包括中间的空格，如"姓 名" -> "姓名"）
        name = name.replace(' ', '')
        
        # 去除首尾空白
        name = name.strip()
        
        return name
    
    def _detect_header_row(self, sheet: Worksheet) -> Tuple[int, Dict[str, int]]:
        """
        检测表头行和列映射（支持多级表头）
        
        Args:
            sheet: 工作表对象
            
        Returns:
            (表头行号, {字段名: 列号})
        """
        # 尝试前10行，检查哪一行包含最多的已知列名
        max_header_row = min(10, sheet.max_row)
        
        best_row = 1
        best_columns = {}
        best_score = 0
        
        # 关键字段权重
        field_weights = {
            'name': 10,           # 姓名 - 最重要
            'sequence': 5,         # 序号
            'branch_name': 8,      # 所属党支部
            'branch_sequence': 5,  # 支部序号
            'position_salary': 3,  # 岗位工资
            'rank_salary': 3,      # 薪级工资
            'housing_fund': 3,     # 住房公积金
            'payment_base': 4,     # 缴费基数
            'monthly_fee': 4,      # 每月应缴党费
        }
        
        # 检查每一行
        for row in range(1, max_header_row + 1):
            columns_found = {}
            row_score = 0
            
            for col in range(1, sheet.max_column + 1):
                cell_value = self._get_merged_cell_value(sheet, row, col)
                if cell_value is None:
                    continue
                
                normalized_name = self._normalize_column_name(cell_value)
                
                # 检查是否匹配任何已知列名
                if normalized_name in self.reverse_mapping:
                    field_name = self.reverse_mapping[normalized_name]
                    columns_found[field_name] = col
                    row_score += field_weights.get(field_name, 1)
            
            # 检查是否有更好的结果
            if row_score > best_score:
                best_score = row_score
                best_row = row
                best_columns = columns_found
        
        # 如果找到的列不够，尝试使用"备用检测方法"
        # 检查是否有"姓名"或"序号"在后面的行
        if best_score < 5:
            # 尝试从第2行开始，逐行检查是否有连续的"序号"、"姓名"等
            for row in range(2, min(max_header_row + 1, sheet.max_row)):
                columns_found = {}
                row_score = 0
                
                for col in range(1, sheet.max_column + 1):
                    cell_value = self._get_merged_cell_value(sheet, row, col)
                    if cell_value is None:
                        continue
                    
                    normalized_name = self._normalize_column_name(cell_value)
                    
                    # 更宽松的匹配
                    # 检查是否包含关键字
                    if '姓名' in normalized_name:
                        columns_found['name'] = col
                        row_score += 10
                    elif '序号' in normalized_name and '支部' not in normalized_name:
                        columns_found['sequence'] = col
                        row_score += 5
                    elif '支部序号' in normalized_name:
                        columns_found['branch_sequence'] = col
                        row_score += 5
                    elif '党支部' in normalized_name or '所属支部' in normalized_name:
                        columns_found['branch_name'] = col
                        row_score += 8
                    elif '岗位工资' in normalized_name:
                        columns_found['position_salary'] = col
                        row_score += 3
                    elif '薪级工资' in normalized_name:
                        columns_found['rank_salary'] = col
                        row_score += 3
                    elif '高定工资' in normalized_name:
                        columns_found['fixed_salary'] = col
                        row_score += 3
                    elif '基础性绩效' in normalized_name:
                        columns_found['basic_performance'] = col
                        row_score += 3
                    elif '住房公积金' in normalized_name:
                        columns_found['housing_fund'] = col
                        row_score += 3
                    elif '医疗保险' in normalized_name:
                        columns_found['medical_insurance'] = col
                        row_score += 3
                    elif '养老保险' in normalized_name:
                        columns_found['pension_insurance'] = col
                        row_score += 3
                    elif '职业年金' in normalized_name:
                        columns_found['occupational_annuity'] = col
                        row_score += 3
                    elif '大额医疗' in normalized_name:
                        columns_found['large_medical'] = col
                        row_score += 3
                    elif '失业保险' in normalized_name:
                        columns_found['unemployment_insurance'] = col
                        row_score += 3
                    elif '个人所得税' in normalized_name:
                        columns_found['personal_income_tax'] = col
                        row_score += 3
                    elif '缴费基数' in normalized_name:
                        columns_found['payment_base'] = col
                        row_score += 4
                    elif '每月应缴党费' in normalized_name or '月党费' in normalized_name:
                        columns_found['monthly_fee'] = col
                        row_score += 4
                
                if row_score > best_score:
                    best_score = row_score
                    best_row = row
                    best_columns = columns_found
        
        print(f"  检测到表头行: 第{best_row}行")
        print(f"  检测到的列映射: {best_columns}")
        print(f"  匹配分数: {best_score}")
        
        return best_row, best_columns
    
    def _get_merged_cell_value(self, sheet: Worksheet, row: int, col: int) -> Any:
        """
        获取合并单元格的值
        
        Args:
            sheet: 工作表对象
            row: 行号
            col: 列号
            
        Returns:
            单元格的值
        """
        # 检查是否是合并单元格
        for merged_range in sheet.merged_cells.ranges:
            if merged_range.min_row <= row <= merged_range.max_row and \
               merged_range.min_col <= col <= merged_range.max_col:
                # 返回合并区域左上角单元格的值
                return sheet.cell(row=merged_range.min_row, column=merged_range.min_col).value
        
        return sheet.cell(row=row, column=col).value
    
    def _extract_branch_name_from_title(self, sheet: Worksheet) -> Optional[str]:
        """
        从表格标题中提取支部名称
        
        Args:
            sheet: 工作表对象
            
        Returns:
            支部名称，如果没有找到返回None
        """
        # 检查前10行的标题
        for row in range(1, min(11, sheet.max_row + 1)):
            for col in range(1, min(20, sheet.max_column + 1)):
                cell_value = sheet.cell(row=row, column=col).value
                if cell_value:
                    text = str(cell_value)
                    # 查找包含"支部"的文本
                    if '支部' in text:
                        # 清理文本
                        text = text.replace('\n', ' ').replace('\r', ' ').strip()
                        # 尝试提取支部名称
                        patterns = [
                            r'([^\s]+?支部)',  # 匹配"XX支部"
                            r'支部[：:]([^\s，。]+)',  # 匹配"支部：XX"
                            r'党支部[：:]([^\s，。]+)',  # 匹配"党支部：XX"
                        ]
                        for pattern in patterns:
                            match = re.search(pattern, text)
                            if match:
                                result = match.group(1).strip()
                                # 排除"党支部"、"各支部"等通用词
                                if result not in ['党支部', '各支部', '支部', '支部名称', '所属支部']:
                                    return result
                        # 如果没有匹配到特定格式，直接使用包含"支部"的文本
                        # 但要排除只是列名的情况
                        if '明细表' in text or '汇总表' in text or '收缴' in text:
                            # 如果是标题，尝试提取
                            for part in text.split():
                                if '支部' in part and part not in ['党支部', '各支部', '支部']:
                                    return part
        return None
    
    def _parse_fee_detail_sheet(self, sheet: Worksheet) -> List[PartyBranch]:
        """
        解析党费收缴明细表（支持总表和子表两种格式）
        
        Args:
            sheet: 工作表对象
            
        Returns:
            党支部列表
        """
        branches_dict = {}
        branches_list = []
        
        # 检测表头行
        header_row, column_mapping = self._detect_header_row(sheet)
        
        # 如果没有找到姓名列，返回空
        if 'name' not in column_mapping and 'sequence' not in column_mapping:
            print("  警告: 未检测到姓名或序号列，无法解析数据")
            return []
        
        # 检查表格类型
        # 如果有 branch_name 或 branch_sequence 列，说明是总表格式
        # 否则，可能是子表格式（整个表格属于同一个支部）
        has_branch_info = ('branch_name' in column_mapping) or ('branch_sequence' in column_mapping)
        
        print(f"  表格类型: {'总表格式（多支部）' if has_branch_info else '子表格式（单支部）'}")
        
        # 如果是子表格式，尝试从标题中提取支部名称
        default_branch_name = None
        if not has_branch_info:
            default_branch_name = self._extract_branch_name_from_title(sheet)
            if not default_branch_name:
                # 使用工作表名称
                default_branch_name = sheet.title
                # 清理工作表名称
                default_branch_name = re.sub(r'党费收缴明细', '', default_branch_name)
                default_branch_name = re.sub(r'党费收缴', '', default_branch_name)
                default_branch_name = default_branch_name.strip()
                if not default_branch_name:
                    default_branch_name = "导入数据"
            print(f"  检测到的支部名称: {default_branch_name}")
            
            # 创建默认支部
            default_branch = PartyBranch(name=default_branch_name, sequence=1)
            branches_dict[default_branch_name] = default_branch
            branches_list.append(default_branch)
        
        # 遍历数据行
        current_branch_sequence = None
        current_branch_name = None
        member_sequence = 0
        
        # 如果是子表格式，设置默认支部名称
        if not has_branch_info:
            current_branch_name = default_branch_name
            current_branch_sequence = 1
        
        for row in range(header_row + 1, sheet.max_row + 1):
            # 获取必要的列
            name_cell = column_mapping.get('name')
            seq_cell = column_mapping.get('sequence')
            
            # 尝试获取姓名
            name_value = None
            if name_cell:
                name_value = self._get_merged_cell_value(sheet, row, name_cell)
            
            # 如果没有姓名列，但有序号列，尝试从下一列获取姓名
            if not name_value and seq_cell:
                # 检查序号列的下一列是否是姓名
                next_col = seq_cell + 1
                if next_col <= sheet.max_column:
                    possible_name = self._get_merged_cell_value(sheet, row, next_col)
                    if possible_name and str(possible_name).strip() not in ["", "合计"]:
                        name_value = possible_name
                        # 动态添加姓名列映射
                        column_mapping['name'] = next_col
            
            # 如果仍然没有姓名，跳过
            if not name_value:
                continue
            
            # 清理姓名
            name_str = str(name_value).strip()
            
            # 跳过合计行和空行
            if name_str == "" or name_str == "合计" or name_str.startswith("合计"):
                continue
            
            # 处理总表格式
            if has_branch_info:
                # 获取支部序号和名称（可能是合并单元格）
                branch_seq_cell = column_mapping.get('branch_sequence')
                branch_name_cell = column_mapping.get('branch_name')
                
                # 获取支部信息
                if branch_seq_cell:
                    branch_seq_value = self._get_merged_cell_value(sheet, row, branch_seq_cell)
                    if branch_seq_value is not None and str(branch_seq_value).strip() != "":
                        try:
                            current_branch_sequence = int(float(branch_seq_value))
                        except (ValueError, TypeError):
                            pass
                
                if branch_name_cell:
                    branch_name_value = self._get_merged_cell_value(sheet, row, branch_name_cell)
                    if branch_name_value is not None and str(branch_name_value).strip() != "":
                        current_branch_name = str(branch_name_value).strip()
                
                # 如果没有支部信息，跳过（在总表格式下）
                if not current_branch_name:
                    continue
            
            # 获取或创建支部
            if current_branch_name not in branches_dict:
                branch = PartyBranch(
                    name=current_branch_name,
                    sequence=current_branch_sequence if current_branch_sequence else len(branches_list) + 1
                )
                branches_dict[current_branch_name] = branch
                branches_list.append(branch)
            else:
                branch = branches_dict[current_branch_name]
            
            # 解析党员信息
            member = PartyMember(
                name=str(name_value).strip(),
                branch_name=current_branch_name,
                branch_sequence=branch.sequence
            )
            
            # 获取党员序号
            if seq_cell:
                seq_value = self._get_merged_cell_value(sheet, row, seq_cell)
                if seq_value is not None:
                    try:
                        member.sequence = int(float(seq_value))
                    except (ValueError, TypeError):
                        member_sequence += 1
                        member.sequence = member_sequence
            else:
                member_sequence += 1
                member.sequence = member_sequence
            
            # 解析工资信息
            salary_info = SalaryInfo()
            deduction_info = DeductionInfo()
            
            # 工资项目
            if 'position_salary' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['position_salary'])
                if value is not None:
                    try:
                        salary_info.position_salary = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'rank_salary' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['rank_salary'])
                if value is not None:
                    try:
                        salary_info.rank_salary = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'fixed_salary' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['fixed_salary'])
                if value is not None:
                    try:
                        salary_info.fixed_salary = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'basic_performance' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['basic_performance'])
                if value is not None:
                    try:
                        salary_info.basic_performance = float(value)
                    except (ValueError, TypeError):
                        pass
            
            # 扣款项目
            if 'housing_fund' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['housing_fund'])
                if value is not None:
                    try:
                        deduction_info.housing_fund = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'medical_insurance' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['medical_insurance'])
                if value is not None:
                    try:
                        deduction_info.medical_insurance = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'pension_insurance' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['pension_insurance'])
                if value is not None:
                    try:
                        deduction_info.pension_insurance = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'occupational_annuity' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['occupational_annuity'])
                if value is not None:
                    try:
                        deduction_info.occupational_annuity = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'large_medical' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['large_medical'])
                if value is not None:
                    try:
                        deduction_info.large_medical = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'unemployment_insurance' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['unemployment_insurance'])
                if value is not None:
                    try:
                        deduction_info.unemployment_insurance = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'personal_income_tax' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['personal_income_tax'])
                if value is not None:
                    try:
                        deduction_info.personal_income_tax = float(value)
                    except (ValueError, TypeError):
                        pass
            
            member.salary_info = salary_info
            member.deduction_info = deduction_info
            
            # 尝试读取已有的缴费基数和党费
            if 'payment_base' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['payment_base'])
                if value is not None:
                    try:
                        member.payment_base = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'monthly_fee' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['monthly_fee'])
                if value is not None:
                    try:
                        member.monthly_fee = float(value)
                    except (ValueError, TypeError):
                        pass
            
            # 如果没有读取到缴费基数或党费，自动计算
            if member.payment_base == 0 or member.monthly_fee == 0:
                FeeCalculator.calculate_member_fee(member)
            
            # 添加到支部
            branch.members.append(member)
        
        # 重新分配党员序号（如果需要）
        for branch in branches_list:
            for i, member in enumerate(branch.members, 1):
                if member.sequence == 0:
                    member.sequence = i
        
        return branches_list
    
    def _parse_branch_sheet(self, sheet: Worksheet, branch_name: str = None) -> Optional[PartyBranch]:
        """
        解析单个支部的党费收缴子表
        
        Args:
            sheet: 工作表对象
            branch_name: 支部名称（如果已知）
            
        Returns:
            党支部对象
        """
        # 首先尝试从标题中提取支部名称
        if not branch_name:
            # 检查前5行的标题
            for row in range(1, min(6, sheet.max_row + 1)):
                for col in range(1, min(10, sheet.max_column + 1)):
                    cell_value = sheet.cell(row=row, column=col).value
                    if cell_value:
                        text = str(cell_value)
                        # 查找包含"支部"的文本
                        if '支部' in text:
                            # 尝试提取支部名称
                            patterns = [
                                r'(.+?支部)',  # 匹配"XX支部"
                                r'支部：(.+?)[，。\n]',  # 匹配"支部：XX"
                                r'党支部[：:](.+?)[，。\n]',  # 匹配"党支部：XX"
                            ]
                            for pattern in patterns:
                                match = re.search(pattern, text)
                                if match:
                                    branch_name = match.group(1).strip()
                                    break
                            if not branch_name:
                                # 如果没有匹配到，直接使用包含"支部"的文本
                                branch_name = text
                            break
                if branch_name:
                    break
        
        # 如果还是没有找到支部名称，使用工作表名称
        if not branch_name:
            branch_name = sheet.title
            # 清理工作表名称
            branch_name = re.sub(r'党费收缴明细', '', branch_name)
            branch_name = re.sub(r'党费收缴', '', branch_name)
            branch_name = branch_name.strip()
        
        if not branch_name:
            branch_name = "未知支部"
        
        # 创建支部对象
        branch = PartyBranch(name=branch_name, sequence=1)
        
        # 检测表头行
        header_row, column_mapping = self._detect_header_row(sheet)
        
        print(f"  解析支部: {branch_name}")
        print(f"  检测到表头行: 第{header_row}行")
        
        # 遍历数据行
        member_sequence = 0
        
        for row in range(header_row + 1, sheet.max_row + 1):
            # 获取姓名
            name_cell = column_mapping.get('name') or column_mapping.get('sequence')
            
            if 'name' in column_mapping:
                name_value = self._get_merged_cell_value(sheet, row, column_mapping['name'])
            else:
                # 如果没有姓名字段，尝试从第二列获取
                name_value = self._get_merged_cell_value(sheet, row, 2)
            
            # 如果姓名为空或为"合计"，跳过
            if not name_value or str(name_value).strip() == "" or str(name_value).strip() == "合计":
                continue
            
            # 解析党员信息
            member = PartyMember(
                name=str(name_value).strip(),
                branch_name=branch_name,
                branch_sequence=1
            )
            
            # 获取党员序号
            if 'sequence' in column_mapping:
                seq_value = self._get_merged_cell_value(sheet, row, column_mapping['sequence'])
                if seq_value is not None:
                    try:
                        member.sequence = int(float(seq_value))
                    except (ValueError, TypeError):
                        member_sequence += 1
                        member.sequence = member_sequence
            else:
                member_sequence += 1
                member.sequence = member_sequence
            
            # 解析工资和扣款信息（与总表解析相同）
            salary_info = SalaryInfo()
            deduction_info = DeductionInfo()
            
            # 工资项目
            for field in ['position_salary', 'rank_salary', 'fixed_salary', 'basic_performance']:
                if field in column_mapping:
                    value = self._get_merged_cell_value(sheet, row, column_mapping[field])
                    if value is not None:
                        try:
                            setattr(salary_info, field, float(value))
                        except (ValueError, TypeError):
                            pass
            
            # 扣款项目
            for field in ['housing_fund', 'medical_insurance', 'pension_insurance', 
                         'occupational_annuity', 'large_medical', 'unemployment_insurance', 
                         'personal_income_tax']:
                if field in column_mapping:
                    value = self._get_merged_cell_value(sheet, row, column_mapping[field])
                    if value is not None:
                        try:
                            setattr(deduction_info, field, float(value))
                        except (ValueError, TypeError):
                            pass
            
            member.salary_info = salary_info
            member.deduction_info = deduction_info
            
            # 尝试读取已有的缴费基数和党费
            if 'payment_base' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['payment_base'])
                if value is not None:
                    try:
                        member.payment_base = float(value)
                    except (ValueError, TypeError):
                        pass
            
            if 'monthly_fee' in column_mapping:
                value = self._get_merged_cell_value(sheet, row, column_mapping['monthly_fee'])
                if value is not None:
                    try:
                        member.monthly_fee = float(value)
                    except (ValueError, TypeError):
                        pass
            
            # 如果没有读取到缴费基数或党费，自动计算
            if member.payment_base == 0 or member.monthly_fee == 0:
                FeeCalculator.calculate_member_fee(member)
            
            # 添加到支部
            branch.members.append(member)
        
        # 重新分配党员序号
        for i, member in enumerate(branch.members, 1):
            if member.sequence == 0:
                member.sequence = i
        
        return branch if branch.members else None
    
    def import_from_file(self, filepath: str, sheet_type: str = 'auto') -> ImportResult:
        """
        从Excel文件导入数据
        
        Args:
            filepath: Excel文件路径
            sheet_type: 表格类型
                - 'auto': 自动检测
                - 'detail': 党费收缴明细表（总表）
                - 'branch': 各支部党费收缴子表（可能包含多个工作表）
                - 'summary': 党费汇总表
            
        Returns:
            导入结果
        """
        result = ImportResult()
        
        # 检查文件是否存在
        if not os.path.exists(filepath):
            result.errors.append(f"文件不存在: {filepath}")
            return result
        
        # 检查文件扩展名
        if not filepath.endswith(('.xlsx', '.xls')):
            result.errors.append(f"不支持的文件格式: {filepath}")
            result.errors.append("仅支持 .xlsx 和 .xls 格式的Excel文件")
            return result
        
        try:
            # 加载工作簿
            wb = load_workbook(filepath, data_only=True)
            result.warnings.append(f"成功加载工作簿: {os.path.basename(filepath)}")
            result.warnings.append(f"工作表数量: {len(wb.sheetnames)}")
            
            # 根据类型解析
            if sheet_type == 'detail' or (sheet_type == 'auto' and len(wb.sheetnames) == 1):
                # 解析总表
                sheet = wb.active
                result.warnings.append(f"解析工作表: {sheet.title}")
                
                branches = self._parse_fee_detail_sheet(sheet)
                result.branches = branches
                result.total_branches = len(branches)
                result.total_members = sum(len(b.members) for b in branches)
                
                if result.total_members > 0:
                    result.success = True
                else:
                    result.errors.append("未找到任何党员数据")
            
            elif sheet_type == 'branch' or (sheet_type == 'auto' and len(wb.sheetnames) > 1):
                # 解析多个工作表（每个工作表是一个支部）
                all_branches = []
                
                for sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                    result.warnings.append(f"解析工作表: {sheet_name}")
                    
                    branch = self._parse_branch_sheet(sheet)
                    if branch and branch.members:
                        # 分配支部序号
                        branch.sequence = len(all_branches) + 1
                        # 更新党员的支部序号
                        for member in branch.members:
                            member.branch_sequence = branch.sequence
                        all_branches.append(branch)
                
                result.branches = all_branches
                result.total_branches = len(all_branches)
                result.total_members = sum(len(b.members) for b in all_branches)
                
                if result.total_members > 0:
                    result.success = True
                else:
                    result.errors.append("未找到任何党员数据")
            
            elif sheet_type == 'summary':
                # 汇总表（暂时只提示）
                result.warnings.append("检测到汇总表类型")
                result.warnings.append("汇总表目前仅用于参考，不会导入党员数据")
                result.success = True
            
            # 关闭工作簿
            wb.close()
            
        except Exception as e:
            result.errors.append(f"导入失败: {str(e)}")
            import traceback
            result.errors.append(traceback.format_exc())
        
        return result
    
    def import_from_files(self, filepaths: List[str]) -> ImportResult:
        """
        从多个Excel文件导入数据
        
        Args:
            filepaths: Excel文件路径列表
            
        Returns:
            导入结果
        """
        result = ImportResult()
        all_branches_dict = {}
        
        for filepath in filepaths:
            file_result = self.import_from_file(filepath)
            
            # 合并警告和错误
            result.warnings.extend(file_result.warnings)
            result.errors.extend(file_result.errors)
            
            # 合并支部数据
            for branch in file_result.branches:
                # 检查是否已存在同名支部
                if branch.name in all_branches_dict:
                    # 合并党员
                    existing_branch = all_branches_dict[branch.name]
                    # 重新分配序号
                    start_seq = len(existing_branch.members) + 1
                    for i, member in enumerate(branch.members):
                        member.sequence = start_seq + i
                        member.branch_sequence = existing_branch.sequence
                    existing_branch.members.extend(branch.members)
                else:
                    # 新支部
                    branch.sequence = len(all_branches_dict) + 1
                    for member in branch.members:
                        member.branch_sequence = branch.sequence
                    all_branches_dict[branch.name] = branch
        
        # 转换为列表
        result.branches = list(all_branches_dict.values())
        # 按支部序号排序
        result.branches.sort(key=lambda b: b.sequence)
        
        result.total_branches = len(result.branches)
        result.total_members = sum(len(b.members) for b in result.branches)
        
        if result.total_members > 0:
            result.success = True
        
        return result
    
    def detect_file_type(self, filepath: str) -> str:
        """
        检测Excel文件类型
        
        Args:
            filepath: Excel文件路径
            
        Returns:
            文件类型: 'detail', 'branch', 'summary', 'unknown'
        """
        if not os.path.exists(filepath):
            return 'unknown'
        
        # 根据文件名判断
        filename = os.path.basename(filepath).lower()
        
        if '明细表' in filename or '明细' in filename and '子表' not in filename:
            return 'detail'
        
        if '子表' in filename or '支部' in filename and '汇总' not in filename:
            return 'branch'
        
        if '汇总' in filename:
            return 'summary'
        
        # 根据工作表数量判断
        try:
            wb = load_workbook(filepath, read_only=True)
            sheet_count = len(wb.sheetnames)
            wb.close()
            
            if sheet_count > 1:
                return 'branch'
            else:
                return 'detail'
        except:
            pass
        
        return 'unknown'
