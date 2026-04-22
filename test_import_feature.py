# 测试导入功能
import os
import sys
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# 添加当前目录到路径
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from config import OUTPUT_DIR, DATA_DIR
from data_models import PartyMember, PartyBranch, SalaryInfo, DeductionInfo
from excel_generator import ExcelGenerator
from excel_importer import ExcelImporter
from member_manager import MemberManager


def create_test_excel_file():
    """创建测试用的Excel文件"""
    print("\n" + "=" * 60)
    print("创建测试Excel文件")
    print("=" * 60)
    
    # 创建测试工作簿（明细表格式）
    wb = Workbook()
    ws = wb.active
    ws.title = "党费收缴明细表"
    
    # 写入标题
    ws['A1'] = "2026年2月党费收缴明细表"
    ws.merge_cells('A1:N1')
    
    # 写入表头
    headers = [
        '支部序号', '所属党支部', '序号', '姓名',
        '岗位工资', '薪级工资', '高定工资', '基础性绩效',
        '住房公积金', '医疗保险', '养老保险',
        '缴费基数', '每月应缴党费'
    ]
    
    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=3, column=col_idx, value=header)
    
    # 写入测试数据
    test_data = [
        # 支部1：教工第一党支部
        (1, '教工第一党支部', 1, '张三', 5000, 3000, 0, 2000, 800, 200, 400, 0, 0),
        (1, '教工第一党支部', 2, '李四', 6000, 4000, 0, 2500, 900, 250, 500, 0, 0),
        # 支部2：本科生第一党支部
        (2, '本科生第一党支部', 1, '王五', 3000, 2000, 0, 1000, 500, 100, 200, 0, 0),
        (2, '本科生第一党支部', 2, '赵六', 4000, 2500, 0, 1500, 600, 150, 300, 0, 0),
        (2, '本科生第一党支部', 3, '钱七', 3500, 2200, 0, 1200, 550, 120, 250, 0, 0),
    ]
    
    for row_idx, row_data in enumerate(test_data, 4):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # 保存文件
    test_file_path = os.path.join(OUTPUT_DIR, '测试_党费收缴明细表.xlsx')
    
    # 确保输出目录存在
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    
    wb.save(test_file_path)
    print(f"[OK] 已创建测试文件: {test_file_path}")
    print(f"     包含: 2个支部, 5名党员")
    
    return test_file_path


def create_test_branch_excel_file():
    """创建测试用的子表Excel文件（多个工作表）"""
    print("\n" + "=" * 60)
    print("创建测试子表Excel文件（多工作表）")
    print("=" * 60)
    
    wb = Workbook()
    
    # 第一个工作表：教工第一党支部
    ws1 = wb.active
    ws1.title = "教工第一党支部"
    
    ws1['A1'] = "教工第一党支部 - 党费收缴明细表"
    ws1.merge_cells('A1:L1')
    
    headers1 = ['序号', '姓名', '岗位工资', '薪级工资', '高定工资', '基础性绩效',
                '住房公积金', '医疗保险', '养老保险', '缴费基数', '每月应缴党费']
    
    for col_idx, header in enumerate(headers1, 1):
        ws1.cell(row=3, column=col_idx, value=header)
    
    # 数据
    ws1.cell(row=4, column=1, value=1)
    ws1.cell(row=4, column=2, value='张三')
    ws1.cell(row=4, column=3, value=5000)
    ws1.cell(row=4, column=4, value=3000)
    ws1.cell(row=4, column=5, value=0)
    ws1.cell(row=4, column=6, value=2000)
    ws1.cell(row=4, column=7, value=800)
    ws1.cell(row=4, column=8, value=200)
    ws1.cell(row=4, column=9, value=400)
    
    ws1.cell(row=5, column=1, value=2)
    ws1.cell(row=5, column=2, value='李四')
    ws1.cell(row=5, column=3, value=6000)
    ws1.cell(row=5, column=4, value=4000)
    ws1.cell(row=5, column=5, value=0)
    ws1.cell(row=5, column=6, value=2500)
    ws1.cell(row=5, column=7, value=900)
    ws1.cell(row=5, column=8, value=250)
    ws1.cell(row=5, column=9, value=500)
    
    # 第二个工作表：本科生第一党支部
    ws2 = wb.create_sheet("本科生第一党支部")
    
    ws2['A1'] = "本科生第一党支部 - 党费收缴明细表"
    ws2.merge_cells('A1:L1')
    
    for col_idx, header in enumerate(headers1, 1):
        ws2.cell(row=3, column=col_idx, value=header)
    
    # 数据
    ws2.cell(row=4, column=1, value=1)
    ws2.cell(row=4, column=2, value='王五')
    ws2.cell(row=4, column=3, value=3000)
    ws2.cell(row=4, column=4, value=2000)
    ws2.cell(row=4, column=5, value=0)
    ws2.cell(row=4, column=6, value=1000)
    ws2.cell(row=4, column=7, value=500)
    ws2.cell(row=4, column=8, value=100)
    ws2.cell(row=4, column=9, value=200)
    
    ws2.cell(row=5, column=1, value=2)
    ws2.cell(row=5, column=2, value='赵六')
    ws2.cell(row=5, column=3, value=4000)
    ws2.cell(row=5, column=4, value=2500)
    ws2.cell(row=5, column=5, value=0)
    ws2.cell(row=5, column=6, value=1500)
    ws2.cell(row=5, column=7, value=600)
    ws2.cell(row=5, column=8, value=150)
    ws2.cell(row=5, column=9, value=300)
    
    # 保存文件
    test_file_path = os.path.join(OUTPUT_DIR, '测试_党费收缴子表.xlsx')
    wb.save(test_file_path)
    
    print(f"[OK] 已创建测试文件: {test_file_path}")
    print(f"     包含: 2个工作表(支部), 4名党员")
    
    return test_file_path


def test_excel_importer_detail(file_path):
    """测试Excel明细表导入"""
    print("\n" + "=" * 60)
    print("测试Excel明细表导入")
    print("=" * 60)
    
    importer = ExcelImporter()
    
    # 检测文件类型
    file_type = importer.detect_file_type(file_path)
    print(f"检测到的文件类型: {file_type}")
    
    # 导入数据
    result = importer.import_from_file(file_path, sheet_type='detail')
    
    print(f"\n导入结果:")
    print(f"  成功: {result.success}")
    print(f"  支部数量: {result.total_branches}")
    print(f"  党员数量: {result.total_members}")
    
    if result.warnings:
        print(f"\n警告信息:")
        for w in result.warnings:
            print(f"  - {w}")
    
    if result.errors:
        print(f"\n错误信息:")
        for e in result.errors:
            print(f"  - {e}")
    
    if result.success:
        print(f"\n导入的支部详情:")
        print("-" * 60)
        for branch in result.branches:
            print(f"\n支部: {branch.name} (序号: {branch.sequence})")
            print(f"  党员数量: {len(branch.members)}")
            for member in branch.members:
                print(f"    - {member.name}: 缴费基数={member.payment_base:.2f}, 党费={member.monthly_fee:.2f}")
    
    return result


def test_excel_importer_branch(file_path):
    """测试Excel子表导入"""
    print("\n" + "=" * 60)
    print("测试Excel子表导入（多工作表）")
    print("=" * 60)
    
    importer = ExcelImporter()
    
    # 检测文件类型
    file_type = importer.detect_file_type(file_path)
    print(f"检测到的文件类型: {file_type}")
    
    # 导入数据
    result = importer.import_from_file(file_path, sheet_type='branch')
    
    print(f"\n导入结果:")
    print(f"  成功: {result.success}")
    print(f"  支部数量: {result.total_branches}")
    print(f"  党员数量: {result.total_members}")
    
    if result.warnings:
        print(f"\n警告信息:")
        for w in result.warnings:
            print(f"  - {w}")
    
    if result.errors:
        print(f"\n错误信息:")
        for e in result.errors:
            print(f"  - {e}")
    
    if result.success:
        print(f"\n导入的支部详情:")
        print("-" * 60)
        for branch in result.branches:
            print(f"\n支部: {branch.name} (序号: {branch.sequence})")
            print(f"  党员数量: {len(branch.members)}")
            for member in branch.members:
                print(f"    - {member.name}: 缴费基数={member.payment_base:.2f}, 党费={member.monthly_fee:.2f}")
    
    return result


def test_directory_structure():
    """测试目录结构（报表按月份组织）"""
    print("\n" + "=" * 60)
    print("测试报表目录结构")
    print("=" * 60)
    
    # 创建测试数据
    manager = MemberManager()
    
    # 添加测试支部
    branch1 = PartyBranch(name="教工第一党支部", sequence=1)
    branch2 = PartyBranch(name="本科生第一党支部", sequence=2)
    
    # 添加测试党员
    member1 = PartyMember(name="张三", branch_name="教工第一党支部", branch_sequence=1, sequence=1)
    member1.salary_info.position_salary = 5000
    member1.salary_info.rank_salary = 3000
    member1.salary_info.basic_performance = 2000
    member1.deduction_info.housing_fund = 800
    member1.deduction_info.medical_insurance = 200
    member1.deduction_info.pension_insurance = 400
    
    from fee_calculator import FeeCalculator
    FeeCalculator.calculate_member_fee(member1)
    
    member2 = PartyMember(name="李四", branch_name="教工第一党支部", branch_sequence=1, sequence=2)
    member2.salary_info.position_salary = 6000
    member2.salary_info.rank_salary = 4000
    member2.salary_info.basic_performance = 2500
    member2.deduction_info.housing_fund = 900
    member2.deduction_info.medical_insurance = 250
    member2.deduction_info.pension_insurance = 500
    FeeCalculator.calculate_member_fee(member2)
    
    branch1.members = [member1, member2]
    
    member3 = PartyMember(name="王五", branch_name="本科生第一党支部", branch_sequence=2, sequence=1)
    member3.salary_info.position_salary = 3000
    member3.salary_info.rank_salary = 2000
    member3.salary_info.basic_performance = 1000
    member3.deduction_info.housing_fund = 500
    member3.deduction_info.medical_insurance = 100
    member3.deduction_info.pension_insurance = 200
    FeeCalculator.calculate_member_fee(member3)
    
    branch2.members = [member3]
    
    manager.add_branch(branch1)
    manager.add_branch(branch2)
    
    # 生成报表
    generator = ExcelGenerator()
    
    print("\n生成报表...")
    
    # 生成明细表
    detail_path = generator.generate_fee_detail_sheet(
        manager.get_all_branches(),
        year=2026,
        month=3
    )
    print(f"[OK] 明细表: {detail_path}")
    
    # 生成子表
    branch_paths = generator.generate_all_branch_sheets(
        manager.get_all_branches(),
        year=2026,
        month=3
    )
    print(f"[OK] 子表: {len(branch_paths)} 个文件")
    for path in branch_paths:
        print(f"     - {path}")
    
    # 生成汇总表
    summary_path = generator.generate_summary_sheet(
        manager.get_all_branches(),
        year=2026,
        month=3
    )
    print(f"[OK] 汇总表: {summary_path}")
    
    # 显示目录结构
    print("\n目录结构:")
    print("-" * 60)
    
    month_dir = os.path.join(OUTPUT_DIR, "2026年3月")
    if os.path.exists(month_dir):
        for root, dirs, files in os.walk(month_dir):
            level = root.replace(month_dir, '').count(os.sep)
            indent = ' ' * 2 * level
            print(f"{indent}{os.path.basename(root)}/")
            subindent = ' ' * 2 * (level + 1)
            for file in files:
                print(f"{subindent}{file}")
    
    return True


def main():
    """主测试函数"""
    print("\n" + "#" * 60)
    print("# 党费收取系统 - 导入功能测试")
    print("#" * 60)
    
    all_passed = True
    
    try:
        # 1. 创建测试文件
        test_detail_file = create_test_excel_file()
        test_branch_file = create_test_branch_excel_file()
        
        # 2. 测试明细表导入
        result1 = test_excel_importer_detail(test_detail_file)
        if not result1.success:
            all_passed = False
        
        # 3. 测试子表导入
        result2 = test_excel_importer_branch(test_branch_file)
        if not result2.success:
            all_passed = False
        
        # 4. 测试目录结构
        test_directory_structure()
        
    except Exception as e:
        print(f"\n[FAIL] 测试过程中出现异常: {e}")
        import traceback
        traceback.print_exc()
        all_passed = False
    
    print("\n" + "#" * 60)
    if all_passed:
        print("# 所有测试通过！")
    else:
        print("# 部分测试失败，请检查错误信息")
    print("#" * 60)


if __name__ == "__main__":
    main()
